import json
import re
from copy import deepcopy
from pathlib import Path

from openpyxl import load_workbook


SOURCE_DIR = Path("/Users/book/Documents/AI SME /outputs/fwd_smart_inventory_2026-05-11")
SOURCE_JSON = SOURCE_DIR / "fwd_smart_inventory_data.json"
SOURCE_XLSX = SOURCE_DIR / "fwd_smart_inventory.xlsx"
OUTPUT_DIR = Path(__file__).resolve().parents[1] / "data" / "fwd"
OUTPUT_JSON = OUTPUT_DIR / "fwd-smart-inventory-2026-05-11.json"
OUTPUT_XLSX = OUTPUT_DIR / "fwd-smart-inventory-2026-05-11.xlsx"

SENSITIVE_KEY_RE = re.compile(
    r"(^|_)(dob|date_of_birth|birth_date|hkid|id_no|identity|phone|mobile|email|address)(_|$)",
    re.IGNORECASE,
)
DOB_TEXT_RE = re.compile(
    r"((?:parent|insured|proposer|policyholder|applicant|客戶|受保人|投保人)?\s*DOB\s*)\d{4}-\d{2}-\d{2}",
    re.IGNORECASE,
)


def redact_text(value: str) -> str:
    return DOB_TEXT_RE.sub(r"\1REDACTED", value)


def sanitize(value, key: str | None = None):
    if key and SENSITIVE_KEY_RE.search(key):
        return "REDACTED" if value not in (None, "") else value
    if isinstance(value, dict):
        return {k: sanitize(v, k) for k, v in value.items()}
    if isinstance(value, list):
        return [sanitize(item, key) for item in value]
    if isinstance(value, str):
        return redact_text(value)
    return value


def sanitize_workbook():
    workbook = load_workbook(SOURCE_XLSX)
    for sheet in workbook.worksheets:
        headers = {}
        for cell in sheet[1]:
            if cell.value:
                headers[cell.column] = str(cell.value)

        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                header = headers.get(cell.column, "")
                if header and SENSITIVE_KEY_RE.search(header):
                    if cell.value not in (None, ""):
                        cell.value = "REDACTED"
                elif isinstance(cell.value, str):
                    cell.value = redact_text(cell.value)

    workbook.save(OUTPUT_XLSX)


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    data = json.loads(SOURCE_JSON.read_text(encoding="utf-8"))
    sanitized = sanitize(deepcopy(data))
    sanitized["_metadata"] = {
        "source_folder": str(SOURCE_DIR),
        "source_json": SOURCE_JSON.name,
        "source_xlsx": SOURCE_XLSX.name,
        "sanitized": True,
        "redaction_rule": "DOB/HKID/phone/email/address-style fields and DOB text snippets redacted before public GitHub commit.",
    }
    OUTPUT_JSON.write_text(json.dumps(sanitized, ensure_ascii=False, indent=2), encoding="utf-8")
    sanitize_workbook()
    print(f"Wrote {OUTPUT_JSON}")
    print(f"Wrote {OUTPUT_XLSX}")


if __name__ == "__main__":
    main()
